"""
Build the auto-insurance edition of the deck.

Starts from the generic deck (template/MCP_Live_Streaming_Context.pptx), rewrites
every slide for the auto-insurance use case, swaps in insurance diagrams
(diagrams.py) and inserts five new slides. The scenario table is generated from
mcp_streaming.domain.auto_insurance so the deck always matches the demo.

  pip install -r requirements-docs.txt
  python docs/presentation/build_deck.py     # from the repository root
"""

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import diagrams  # noqa: E402
from mcp_streaming.domain import auto_insurance as domain  # noqa: E402

TEMPLATE = os.path.join(HERE, "template", "MCP_Live_Streaming_Context.pptx")
OUTPUT = os.path.join(HERE, "MCP_Live_Auto_Insurance.pptx")

NAVY = RGBColor(0x2A, 0x2A, 0x4A)
TEAL = RGBColor(0x00, 0xD4, 0xAA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF2, 0xF2, 0xF6)
GREY_TEXT = RGBColor(0x55, 0x55, 0x6A)
MONO = "Menlo"


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #

def set_lines(shape, lines, size=14, font=None, color=None):
    """Replace a text frame's paragraphs, mirroring the template's 14pt / 6pt-after style."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        for j, segment in enumerate(line.split("\n")):  # "\n" inside a line = soft line break
            if j:
                p.add_line_break()
            run = p.add_run()
            run.text = segment
            run.font.size = Pt(size)
            if font:
                run.font.name = font
            if color is not None:
                run.font.color.rgb = color


def title_of(slide):
    return slide.shapes.title


def body_placeholders(slide):
    return [s for s in slide.placeholders if s.placeholder_format.idx != 0]


def first_picture(slide):
    return next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)


def replace_picture(slide, image_path, box=None):
    """Swap the slide's picture for a new image, fitted (aspect preserved) and centred in the same box."""
    old = first_picture(slide)
    left, top, width, height = box or (old.left, old.top, old.width, old.height)
    old._element.getparent().remove(old._element)
    from PIL import Image
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(width / iw, height / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(image_path, left + (width - w) // 2, top + (height - h) // 2, w, h)


def add_slide(prs, layout_name, position):
    layout = next(l for l in prs.slide_layouts if l.name == layout_name)
    slide = prs.slides.add_slide(layout)
    ids = prs.slides._sldIdLst
    el = ids[-1]
    ids.remove(el)
    ids.insert(position, el)
    return slide


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_table(slide, rows, col_widths_in, top_in=1.15, font_pt=10, header_pt=11):
    left, width = Inches(0.3), Inches(sum(col_widths_in))
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, Inches(top_in), width, Inches(0.3 * len(rows)))
    table = shape.table
    for i, w in enumerate(col_widths_in):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (ROW_ALT if r % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = value
            run.font.size = Pt(header_pt if r == 0 else font_pt)
            run.font.bold = r == 0 or c == 0
            run.font.color.rgb = WHITE if r == 0 else NAVY
    return table


def add_code_box(slide, lines, left_in, top_in, width_in, height_in, size=10.5):
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.1)
    for i, (text, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.name = MONO
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_caption(slide, text, top_in=5.12):
    box = slide.shapes.add_textbox(Inches(0.3), Inches(top_in), Inches(9.4), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = GREY_TEXT


# --------------------------------------------------------------------------- #
# Build
# --------------------------------------------------------------------------- #

def build():
    img = {fn.__name__: fn() for fn in diagrams.ALL}
    prs = Presentation(TEMPLATE)
    s = list(prs.slides)  # original 13 slides, by position

    # 1 — Title
    set_lines(title_of(s[0]), ["MCP Live for Auto Insurance", "Streaming Claims Context", "to AI Agents"], size=30)
    for p in title_of(s[0]).text_frame.paragraphs:
        p.space_after = Pt(0)
    set_lines(body_placeholders(s[0])[0], [
        "Real-time FNOL, fraud & platform-incident intelligence with Kafka, WebSockets, Cedar & LLMs",
        "github.com/hakohli/mcp-streaming-demo  ·  branch: usecase/auto-insurance"], size=14)
    notes(s[0], "Auto-insurance edition: the same streaming-MCP pattern, applied to an insurer's claims and "
                "telematics platform. All data in the demo is synthetic.")

    # 2 — Architecture
    title_of(s[1]).text = "System Architecture — Claims & Telematics Platform"
    replace_picture(s[1], img["architecture"])
    notes(s[1], "Eight insurance services publish to the insurance-events topic. The MCP server pushes events "
                "to the AI ops agent; on a spike the LLM calls back through Cedar-checked, PII-redacted tools "
                "and consults the insurer's runbooks.")

    # 3 — Stale context
    title_of(s[2]).text = "The Stale Context Problem in Claims"
    set_lines(body_placeholders(s[2])[0], [
        "Claims & telematics platforms are still watched through dashboards and batch reports",
        "2:03 pm hail hits DFW → 2:10 pm FNOL volume 14× → next-morning ops report",
        "Payout failures and fraud rings often surface days later, in reconciliation",
        "What if claims context came TO the agent, as it happened?"])
    replace_picture(s[2], img["timeline"])

    # 4 — Streaming MCP architecture (generic pull-vs-push image kept)
    title_of(s[3]).text = "Streaming MCP for Insurers"
    set_lines(body_placeholders(s[3])[0], [
        "Standard MCP: agent pulls claim / telematics snapshots (pull)",
        "Streaming MCP: FNOL, telematics, fraud and payout events pushed as they happen",
        "Why Kafka: durable and replayable — already the backbone of many telematics pipelines",
        "Hybrid: live push + on-demand queries on one connection"])

    # 5 — Protocol
    title_of(s[4]).text = "Key Protocol Extension: subscribe"
    set_lines(body_placeholders(s[4])[0], [
        "subscribe / unsubscribe to the insurance-events stream for continuous push",
        "get_anomalies, get_context answer the LLM's tool calls mid-stream",
        "Every tool result is PII-redacted before it reaches the model",
        "Bounded queues (500 events) absorb CAT-surge backpressure"])
    replace_picture(s[4], img["sequence"])

    # 6 — Cedar (generic allow/deny image kept)
    title_of(s[5]).text = "Cedar Authorization for Insurance Roles"
    set_lines(body_placeholders(s[5])[0], [
        "Open-source policy language by AWS: permit / forbid(principal, action, resource)",
        "ai-agent: anomaly summaries + redacted event context",
        "claims-ops: summaries  ·  siu-analyst: full event context",
        "ai-agent is forbidden from the raw claimant-PII stream",
        "Default deny — and forbid always wins"], size=13)

    # 7 — Policy examples
    title_of(s[6]).text = "Cedar Policy Examples"
    set_lines(body_placeholders(s[6])[0], [
        'permit(principal, action == Action::"subscribe", resource == Stream::"insurance-events");',
        'permit(principal, action == Action::"get_context", resource)\n    when { principal.role == "siu-analyst" };',
        'permit(principal, action == Action::"get_anomalies", resource)\n    when { principal.role == "claims-ops" };',
        'forbid(principal, action == Action::"subscribe", resource == Stream::"claimant-pii")\n'
        '    when { principal.role == "ai-agent" };',
    ], size=12.5, font=MONO)
    tf = body_placeholders(s[6])[0].text_frame
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = "Forbid always wins — the AI agent can never read raw claimant PII"
    run.font.size = Pt(14)
    run.font.bold = True

    # 8 — Live demo
    title_of(s[7]).text = "Live Demo: Real-time Claims & Telematics Monitoring"
    set_lines(body_placeholders(s[7])[0], [
        "Simulated claims & telematics platform → Kafka → MCP server → AI agent",
        "7 incidents: outages, vendor failure, bad deploy, CAT surge, fraud ring",
        "Spike detected in < 1 s; the LLM diagnoses in seconds",
        "Output: root cause, business impact, owner, next steps",
        "Try it: make demo  ·  --scenario staged_accident_ring"], size=13)
    replace_picture(s[7], img["pipeline"])

    # 9 — Anomaly detection (generic chart kept)
    title_of(s[8]).text = "Anomaly Detection → AI Diagnosis"
    set_lines(body_placeholders(s[8])[0], [
        "Server keeps a rolling 100-event window; IDs are normalized so errors group",
        "Agent's 20-event window: > 30% errors triggers the LLM",
        "Cooldown caps LLM calls — cost stays flat even during a CAT surge",
        "LLM pulls context + runbooks, then recommends; humans decide"], size=13)

    # 10 — Production patterns (+ insurance-specific)
    set_lines(body_placeholders(s[9])[0], [
        "1. Backpressure: bounded queues, drop slow subscribers",
        "2. Hybrid protocol: streaming AND on-demand on the same WebSocket",
        "3. Kafka consumer groups: same = load balance, different = fan-out",
        "4. Reconnection: exponential backoff + context snapshot",
        "5. Be selective: stream errors & anomalies, batch normal metrics",
        "6. Redact PII before the LLM; keep the agent read-only",
        "7. Separate technical faults from business anomalies — different owners"], size=13)

    # 11 — What it unlocks
    title_of(s[10]).text = "What Streaming MCP Unlocks for Auto Insurers"
    set_lines(body_placeholders(s[10])[0], [
        "FNOL automation — crash events to claims in seconds",
        "Fraud / SIU — ring signals surfaced as they form",
        "CAT response — regional surges detected at onset",
        "UBI telematics — data gaps flagged before pricing",
        "Claims payouts — failures tied to their cause",
        "Platform SRE — autonomous triage across 8 services"], size=13)
    replace_picture(s[10], img["hub"])

    # 12 — Key takeaways
    set_lines(body_placeholders(s[11])[0], [
        "Claims context goes stale in minutes — stream it to the agent",
        "Kafka + WebSocket: a durable, push-based bridge into MCP",
        "LLM + runbooks: root cause, impact, owner and next steps in seconds",
        "Cedar + PII redaction + read-only tools: governance by design",
        "Open source, Python 3.9+, runs offline with make demo"], size=13)
    replace_picture(s[11], img["takeaways"])

    # 13 — Thank you
    set_lines(body_placeholders(s[12])[0], [
        "github.com/hakohli/mcp-streaming-demo  ·  branch: usecase/auto-insurance",
        "Synthetic data only  ·  Kafka · Cedar · WebSockets · LLMs"], size=14)

    # ---- New slides (inserted after "Anomaly Detection", position index 9..13) ----
    pos = 9

    sl = add_slide(prs, "TITLE_ONLY", pos)
    title_of(sl).text = "Auto Insurance Use Cases"
    add_table(sl, [
        ["Use case", "Streaming signal", "What the AI agent does", "Owner"],
        ["Crash-triggered FNOL", "Telematics impact events", "Detects stalled auto-FNOL; routes crash alerts to manual intake", "Claims ops · SRE"],
        ["Catastrophe response", "FNOL volume by region", "Spots the surge at onset; recommends activating the CAT plan", "CAT team"],
        ["Fraud / SIU triage", "Fraud scores, link analysis", "Summarizes ring signals; recommends SIU referral (never auto-deny)", "SIU"],
        ["Claims cycle time", "Estimate & vendor calls", "Catches vendor / estimate failures before backlogs build", "Vendor mgmt"],
        ["Payout reliability", "Payment batches", "Ties payout rejections to the deploy that caused them", "Payments eng"],
        ["UBI pricing integrity", "Trip ingestion, rating", "Flags telematics gaps before premiums use partial data", "Data · SRE"],
    ], [1.8, 1.9, 4.3, 1.4], font_pt=12, header_pt=13)
    notes(sl, "Each row is a stream the insurer already has. The agent's job is the same in every row: notice "
              "fast, explain the impact, route to the right owner, recommend steps.")

    sl = add_slide(prs, "TITLE_ONLY", pos + 1)
    title_of(sl).text = "Incident Scenarios in the Demo"
    rows = [["Scenario", "Type", "Owner", "First recommended step"]]
    for sc in domain.SCENARIOS:
        rows.append([sc["title"], sc["category"], domain.OWNERS[sc["owner"]], sc["remediation"][0]])
    add_table(sl, rows, [2.9, 0.9, 2.0, 3.6], font_pt=9.5, header_pt=10.5)
    add_caption(sl, "Generated from mcp_streaming/domain/auto_insurance.py  ·  run one: "
                    "python -m mcp_streaming.demo --scenario <key>   ·   --ai lets the LLM invent new ones")

    sl = add_slide(prs, "TITLE_ONLY", pos + 2)
    title_of(sl).text = "AI Diagnosis — Example Output"
    sc = domain.SCENARIOS_BY_KEY["payout_batch_rejected"]
    import textwrap

    def field(label, text, color=WHITE):
        first, rest = f"| {label:<12}", "| " + " " * 12
        wrapped = textwrap.wrap(text, 92 - len(first))
        return [(first + wrapped[0], color)] + [(rest + w, color) for w in wrapped[1:]]

    code = [("SPIKE  35% errors in last 20 events  ->  LLM investigating...", TEAL),
            ("+-- AI DIAGNOSIS " + "-" * 60, WHITE)]
    code += field("Severity", "HIGH  (confidence 80%)  [platform]")
    code += field("Root cause", sc["root_cause"])
    code += field("Impact", sc["business_impact"])
    code += field("Owner", domain.OWNERS[sc["owner"]])
    code += [("| Next steps", WHITE)]
    for i, step in enumerate(sc["remediation"], 1):
        code += [(f"|   {i}. {line}" if j == 0 else f"|      {line}", WHITE)
                 for j, line in enumerate(textwrap.wrap(step, 84))]
    code += [("| Recommendations only - claim, coverage and fraud decisions stay with people.", TEAL),
             ("+" + "-" * 77, WHITE)]
    add_code_box(sl, code, 0.3, 1.1, 9.4, 3.9, size=10.5)
    add_caption(sl, "Illustrative output for the payout-rejection scenario. A real LLM writes its own analysis "
                    "after calling get_anomalies, get_context and search_runbooks.")

    sl = add_slide(prs, "TITLE_AND_TWO_COLUMNS", pos + 3)
    title_of(sl).text = "Why It Matters for Insurers"
    left, right = body_placeholders(sl)[:2]
    set_lines(left, [
        "Faster first notice of loss and claim acknowledgement",
        "Lower loss-adjustment expense: vendor and estimate failures caught early",
        "Earlier fraud signals: less leakage, fewer false positives",
        "CAT response triggered in minutes, not the next morning",
        "Fewer delayed payouts and complaints"], size=14)
    set_lines(right, [
        "Who benefits",
        "Claims operations",
        "Special Investigations Unit (SIU)",
        "Catastrophe response",
        "Payments & platform engineering",
        "Customer experience",
        "Compliance & audit"], size=14)
    right.text_frame.paragraphs[0].runs[0].font.bold = True
    notes(sl, "Qualitative benefits — measure your own baseline (time-to-detect, cycle time, LAE, leakage) "
              "before and after a pilot.")

    sl = add_slide(prs, "TITLE_AND_BODY", pos + 4)
    title_of(sl).text = "Governance, Privacy & Compliance"
    set_lines(body_placeholders(sl)[0], [
        "PII (email, phone, SSN, VIN, card) is redacted before any event reaches the LLM",
        "Read-only agent: recommends; never approves or denies claims, changes coverage or accuses claimants",
        "Cedar least-privilege roles; the AI agent is forbidden from the raw claimant-PII stream",
        "Event content is treated as untrusted input (prompt-injection defense)",
        "Model hosting fits your data boundary: Amazon Bedrock in-region, or self-hosted (Ollama / vLLM)",
        "Production also needs: real authentication, audit logging, model-risk governance (e.g. the NAIC model "
        "bulletin on insurers' use of AI), state claim-handling and privacy rules — with your compliance team",
        "This is a reference demo on synthetic data, not a production system"], size=14)

    prs.save(OUTPUT)
    return OUTPUT, len(prs.slides)


if __name__ == "__main__":
    path, n = build()
    print(f"wrote {path} ({n} slides)")
